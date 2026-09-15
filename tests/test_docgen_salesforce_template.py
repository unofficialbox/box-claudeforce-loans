"""Every tag in the Salesforce commitment letter must resolve in the managed package's JSON export."""
import json
import re
import unittest
from pathlib import Path

from docx import Document

from scripts.generate_docgen_templates import SALESFORCE_TEMPLATE, commitment_letter_salesforce

ROOT = Path(__file__).resolve().parents[1]
EXPORT = ROOT / 'sample-data' / 'docgen' / 'LOS_Loan__c_DocGen.json'
TAG = re.compile(r'\{\{([^{}]+)\}\}')


def template_text(doc):
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    section = doc.sections[0]
    parts += [p.text for p in section.header.paragraphs] + [p.text for p in section.footer.paragraphs]
    return '\n'.join(parts)


def resolve(data, path):
    node = data
    for key in path.split('.'):
        if isinstance(node, list):
            node = node[0]
        if not isinstance(node, dict) or key not in node:
            return None
        node = node[key]
    return node


class SalesforceDocGenTemplateTests(unittest.TestCase):
    def setUp(self):
        self.export = json.loads(EXPORT.read_text())
        self.doc = commitment_letter_salesforce()
        self.text = template_text(self.doc)

    def test_export_is_the_managed_package_shape(self):
        self.assertEqual(set(self.export), {'$Organization', '$User', 'LOS_Loan__c'})

    def test_every_tag_resolves_in_the_export(self):
        aliases = {}
        checked = 0
        for raw in TAG.findall(self.text):
            expression = raw.strip()
            if expression in ('else', 'endif', 'endtablerow'):
                continue
            loop = re.fullmatch(r'tablerow (\w+) in ([\w$.]+)', expression)
            if loop:
                alias, path = loop.groups()
                collection = resolve(self.export, path)
                self.assertIsInstance(collection, list, path)
                aliases[alias] = path
                continue
            condition = re.fullmatch(r'if ([\w$.]+) isPresent', expression)
            path = condition.group(1) if condition else expression.split('::', 1)[0].strip()
            self.assertRegex(path, r'^[\w$]+(\.\w+)*$', f'Unsupported tag syntax: {raw!r}')
            head, _, rest = path.partition('.')
            if head in aliases:
                path = aliases[head] + ('.' + rest if rest else '')
            self.assertIsNotNone(resolve(self.export, path), f'Tag {raw!r} does not resolve in {EXPORT.name}')
            checked += 1
        self.assertGreater(checked, 30)

    def test_no_tag_is_split_across_runs(self):
        for paragraph in list(self.doc.paragraphs) + [
            p for table in self.doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs
        ]:
            if '{{' in paragraph.text:
                self.assertEqual(TAG.findall(paragraph.text), [tag for run in paragraph.runs for tag in TAG.findall(run.text)])

    def test_modifiers_use_documented_syntax(self):
        for raw in TAG.findall(self.text):
            if '::' in raw:
                self.assertRegex(raw.split('::', 1)[1].strip(), r'^(optional|format\("(US-Number|EU-Number|dd-mm-yyyy|mm-dd-yyyy|dd-mmm-yy)"\))$', raw)
        self.assertEqual(len(re.findall(r'\{\{ if ', self.text)), len(re.findall(r'\{\{ endif \}\}', self.text)))

    def test_only_constructs_the_managed_package_supplies(self):
        # Verified against the package preview and a direct Box Doc Gen run: `$User` is not
        # resolved, a lookup's child list is not sent, and a missing key prints the tag
        # literally, so every field that can be blank on a loan record must be optional.
        required = {'LOS_Loan__c.Loan_ID__c', 'LOS_Loan__c.Name', 'LOS_Loan__c.Status__c', 'LOS_Loan__c.Underwriting_Notes__c'}
        for raw in TAG.findall(self.text):
            expression = raw.strip()
            self.assertFalse(expression.startswith('$'), raw)
            self.assertNotIn('Borrower_Loans__r', expression)
            self.assertNotIn('tablerow', expression)
            if expression in ('else', 'endif') or expression.startswith('if '):
                continue
            path, _, modifier = (part.strip() for part in expression.partition('::'))
            if path in required or modifier.startswith('format('):
                continue
            self.assertEqual(modifier, 'optional', f'{raw!r} prints literally when the field is null')

    def test_signature_fields_match_the_first_template(self):
        self.assertIn('[[s|1|id:borrower_signature', self.text)
        self.assertIn('[[d|1|id:borrower_signed_date]]', self.text)

    def test_committed_template_matches_generator(self):
        committed = template_text(Document(ROOT / 'output' / 'docgen' / SALESFORCE_TEMPLATE))
        self.assertEqual(committed, self.text)


if __name__ == '__main__':
    unittest.main()
